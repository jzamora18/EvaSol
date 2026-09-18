from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
from PIL import Image


def _reemplazar_parrafo(paragraph, reemplazos: dict) -> bool:
    full_text = "".join(run.text for run in paragraph.runs)

    if "{{IMAGEN}" in full_text or "{{IMAGEN}}" in full_text:
        return True

    if not any(ph in full_text for ph in reemplazos):
        return False

    for placeholder, value in reemplazos.items():
        full_text = full_text.replace(placeholder, value)

    paragraph.runs[0].text = full_text
    for run in paragraph.runs[1:]:
        run.text = ""

    return False


def _nuevo_parrafo_texto(texto, bold=False, pt=10, center=False):
    """Crea un elemento <w:p> con texto para insertar directo en el XML."""
    p = OxmlElement("w:p")

    pPr = OxmlElement("w:pPr")
    if center:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        pPr.append(jc)
    p.append(pPr)

    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if bold:
        b = OxmlElement("w:b")
        rPr.append(b)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(pt * 2))
    rPr.append(sz)
    r.append(rPr)

    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    p.append(r)
    return p


def _insertar_imagenes(doc: Document, placeholder_paragraph, imagenes: list):
    """
    Reemplaza el placeholder con todas las imágenes,
    cada una con su número encima. Todo dentro del mismo doc.
    """
    # Limpiar el placeholder
    for run in placeholder_paragraph.runs:
        run.text = ""

    anchor = placeholder_paragraph._element

    for i, imagen in enumerate(imagenes, start=1):

        # ── Título "Imagen N" ─────────────────────────────────────────────
        title_el = _nuevo_parrafo_texto(f"Imagen {i}", bold=True, pt=10, center=True)
        anchor.addnext(title_el)
        anchor = title_el

        # ── Párrafo de imagen — dentro del mismo doc ──────────────────────
        img_para = doc.add_paragraph()          # se añade al final del doc
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        img_bytes = io.BytesIO()
        imagen.save(img_bytes, format="PNG")
        img_bytes.seek(0)
        img_para.add_run().add_picture(img_bytes, width=Inches(3.5))

        # Mover el elemento XML justo después del título
        anchor.addnext(img_para._element)
        anchor = img_para._element


def _procesar_parrafos(doc: Document, paragraphs, imagenes: list, reemplazos: dict):
    for paragraph in paragraphs:
        if _reemplazar_parrafo(paragraph, reemplazos):
            _insertar_imagenes(doc, paragraph, imagenes)
            return  # solo hay un {{IMAGEN}}


def generar_reporte(
    ruta_plantilla: str,
    imagenes_y_detecciones: list[tuple[Image.Image, list]],
    reemplazos: dict,
) -> bytes:
    doc = Document(ruta_plantilla)

    imagenes = [img for img, _ in imagenes_y_detecciones]

    # ── Cuerpo ────────────────────────────────────────────────────────────────
    _procesar_parrafos(doc, doc.paragraphs, imagenes, reemplazos)

    # ── Tablas ────────────────────────────────────────────────────────────────
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                _procesar_parrafos(doc, cell.paragraphs, imagenes, reemplazos)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()